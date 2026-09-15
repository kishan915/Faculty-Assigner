import io
import pandas as pd
from collections import defaultdict
from flask import Flask, request, send_file, render_template_string
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

# ── Fills ─────────────────────────────────────────────────────────────────────
RED_FILL    = PatternFill("solid", fgColor="FF6B6B")
YELLOW_FILL = PatternFill("solid", fgColor="FFD966")
NO_FILL     = PatternFill(fill_type=None)

# ── HTML ──────────────────────────────────────────────────────────────────────
HTML = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8"/>
  <meta name="viewport" content="width=device-width,initial-scale=1.0"/>
  <title>Faculty Role Assigner</title>
  <style>
    *,*::before,*::after{box-sizing:border-box;margin:0;padding:0}
    body{font-family:'Inter','Segoe UI',system-ui,sans-serif;background:#F0EDE8;
         min-height:100vh;display:flex;align-items:center;justify-content:center;
         padding:2rem 1rem;color:#1C1B19}
    .card{background:#fff;border-radius:16px;padding:2.5rem 2.75rem;
          width:100%;max-width:500px;border:1px solid #E2DED8}
    h1{font-size:22px;font-weight:650;letter-spacing:-.025em;margin-bottom:.4rem}
    .sub{font-size:14px;color:#706E69;margin-bottom:1.75rem;line-height:1.55}
    .rules{background:#F7F5F2;border-radius:10px;padding:.9rem 1rem;
           margin-bottom:1.75rem;border:1px solid #E8E4DF;font-size:12.5px;color:#4A4845}
    .rules b{color:#1C1B19}
    .rules p{margin-bottom:4px;line-height:1.5}
    label{display:block;font-size:13px;font-weight:500;color:#3A3835;margin-bottom:.5rem}
    .drop{border:1.5px dashed #D4D0CA;border-radius:10px;padding:1.5rem 1rem;
          text-align:center;cursor:pointer;background:#FAFAF9;margin-bottom:1rem;
          position:relative;transition:border-color .15s,background .15s}
    .drop:hover,.drop.over{border-color:#1C1B19;background:#F4F2EF}
    .drop input{position:absolute;inset:0;opacity:0;cursor:pointer;width:100%;height:100%}
    .dp{font-size:13.5px;font-weight:500;color:#3A3835;margin-bottom:3px}
    .ds{font-size:12px;color:#9C9A95}
    .chosen{display:none;align-items:center;gap:8px;background:#F0EDE8;
            border-radius:8px;padding:.55rem .85rem;margin-bottom:1rem;
            font-size:13px;color:#3A3835;font-weight:500}
    .btn{width:100%;padding:.7rem 1rem;background:#1C1B19;color:#fff;border:none;
         border-radius:9px;font-size:14px;font-weight:550;cursor:pointer;
         display:flex;align-items:center;justify-content:center;gap:7px;
         transition:background .15s}
    .btn:hover{background:#2E2D2A}
    .btn:disabled{background:#C5C2BC;cursor:not-allowed}
    .credit{
    display:flex;
    align-items:center;
    justify-content:center;
    gap:6px;
    margin-top:.9rem;
    font-size:11.5px;
    color:#9C9A95;
    letter-spacing:.01em;
    }

    .credit-dot{
    width:6px;
    height:6px;
    background:#6B6A67;
    border-radius:50%;
    box-shadow:0 0 5px #6B6A67;
    }
    .spin{display:none;width:15px;height:15px;border:2px solid rgba(255,255,255,.3);
          border-top-color:#fff;border-radius:50%;animation:sp .6s linear infinite}
    @keyframes sp{to{transform:rotate(360deg)}}
    .err{background:#FEF0F0;border:1px solid #FAC5C5;color:#7A1F1F;
         border-radius:9px;padding:.75rem 1rem;font-size:13px;margin-top:1rem}
  </style>
</head>
<body>
<div class="card">
  <h1>Faculty Role Assigner</h1>
  <p class="sub">Upload the faculty workload .xlsx and download the completed assignment file.</p>
  <div class="rules">
    <p><b>Red</b> — No setter found / No faculty available for role</p>
    <p><b>Yellow</b> — Faculty repeated globally in setter role</p>
    <p><b>"not applicable"</b> — Role does not apply for this subject type</p>
  </div>
  <form id="f" action="/generate" method="post" enctype="multipart/form-data">
    <label for="fi">Input file (.xlsx)</label>
    <div class="drop" id="dz">
      <input type="file" name="file" id="fi" accept=".xlsx" required/>
      <p class="dp">Drop your .xlsx here</p>
      <p class="ds">or click to browse</p>
    </div>
    <div class="chosen" id="ch"><span id="fn">file.xlsx</span></div>
    <button class="btn" type="submit" id="sb" disabled>
      <span class="spin" id="sp"></span>
      <span id="bt">Select a file to continue</span>
    </button>
    <div class="credit">
  <span class="credit-dot"></span>
  <span>Created by Kishan Kumar & Kavy Singh Solanki </span>
</div>
  </form>
  {% if error %}<div class="err">{{ error }}</div>{% endif %}
</div>
<script>
const fi=document.getElementById('fi'),dz=document.getElementById('dz'),
  ch=document.getElementById('ch'),fn=document.getElementById('fn'),
  sb=document.getElementById('sb'),bt=document.getElementById('bt'),
  sp=document.getElementById('sp'),frm=document.getElementById('f');
function show(n){dz.style.display='none';ch.style.display='flex';
  fn.textContent=n;sb.disabled=false;bt.textContent='Generate assignment file';}
fi.addEventListener('change',()=>{if(fi.files[0])show(fi.files[0].name);});
dz.addEventListener('dragover',e=>{e.preventDefault();dz.classList.add('over');});
dz.addEventListener('dragleave',()=>dz.classList.remove('over'));
dz.addEventListener('drop',e=>{e.preventDefault();dz.classList.remove('over');
  if(e.dataTransfer.files[0]){fi.files=e.dataTransfer.files;show(e.dataTransfer.files[0].name);}});
frm.addEventListener('submit',()=>{sb.disabled=true;sp.style.display='block';bt.textContent='Generating…';});
</script>
</body>
</html>"""


# ── Subject-level helpers ─────────────────────────────────────────────────────

def detect_subject_type(group):
    if (group["Theory Workload"] == 0).all(): return "Lab"
    if (group["Lab Workload"]    == 0).all(): return "Theory"
    return "Lab+Theory"

def detect_course_type(code):
    c = str(code).strip()
    if c.startswith("2"): return "B. Tech."
    if c.startswith("3"): return "M. Tech."
    return "Unknown"

def extract_semester(code):
    """Skip first digit-run (course prefix); first digit of next run = semester."""
    runs, cur = [], ""
    for ch in str(code).strip():
        if ch.isdigit(): cur += ch
        else:
            if cur: runs.append(cur); cur = ""
    if cur: runs.append(cur)
    return int(runs[1][0]) if len(runs) >= 2 else 0


# ── Core algorithm ────────────────────────────────────────────────────────────

def run_algorithm(df):
    subjects = []

    for (stream, code), grp in df.groupby(["Stream", "Subject code"], sort=False):
        s = {
            "stream":       stream,
            "code":         str(code),
            "name":         grp.iloc[0]["Subject Name"],
            "subject_type": detect_subject_type(grp),
            "course_type":  detect_course_type(code),
            "semester":     extract_semester(code),
        }
        coord = grp[grp["Is Coordinator"] == True]
        s["coordinator"] = coord.iloc[0]["Faculty name"] if not coord.empty else None

        # Sorted pools
        s["all_desc"]    = grp.sort_values("Total Workload", ascending=False).reset_index(drop=True)
        s["theory_desc"] = grp[grp["Theory Workload"] > 0].sort_values("Total Workload", ascending=False).reset_index(drop=True)
        s["lab_desc"]    = grp[grp["Lab Workload"]    > 0].sort_values("Lab Workload",   ascending=False).reset_index(drop=True)

        # Output fields
        s.update({
            "setter1": None, "setter1_red": False, "setter1_yellow": False, "setter1_na": False,
            "setter2": None, "setter2_red": False, "setter2_yellow": False, "setter2_na": False,
            "practical_only": None, "examiner": None, "reserve": None,
            "examiner_na": False,
            "_po_list": [],
        })
        subjects.append(s)

    # ── Global counter (yellow = repeated on a DIFFERENT subject code globally) ──
    # global_counter: faculty_name -> set of unique subject codes assigned as setter
    global_counter = {}   # faculty -> set of codes
    def gc(f, code):
        """Return number of UNIQUE codes this faculty has been assigned to so far."""
        return len(global_counter.get(f, set()))
    def gi(f, code):
        """Record this code for this faculty."""
        if f not in global_counter:
            global_counter[f] = set()
        global_counter[f].add(code)

    # ── Setter assignment ─────────────────────────────────────────────────────
    def assign_setter(s, field, yf, rf, naf, exclude=None):
        """
        Priority tiers — stop at FIRST tier with any eligible candidate.
        Tier 1 (Setter I only): Subject Coordinator
        Tier 2: max Total Workload (all faculty)
        Tier 3: max Total Workload (Theory faculty only)
        Tier 4: max Lab Workload   (Lab faculty only)
        Red if nothing found.  Yellow if repeated globally.
        """
        def pick(pool):
            for _, r in pool.iterrows():
                f = r["Faculty name"]
                if f != exclude: return f
            return None

        def commit(f):
            current_code   = s["code"]
            existing_codes = global_counter.get(f, set())
            # Yellow only if:
            #   1. Faculty already has OTHER codes (not the current one), AND
            #   2. The current code is NEW (first time seeing it)
            # If current_code is already in existing_codes, this is the same subject
            # appearing in a different stream — not a genuine repetition, no yellow.
            is_new_code     = current_code not in existing_codes
            has_other_codes = len(existing_codes - {current_code}) > 0
            s[yf]    = is_new_code and has_other_codes
            s[field] = f
            gi(f, current_code)

        if exclude is None and s["coordinator"]:
            commit(s["coordinator"]); return

        f = pick(s["all_desc"])
        if f: commit(f); return
        f = pick(s["theory_desc"])
        if f: commit(f); return
        f = pick(s["lab_desc"])
        if f: commit(f); return

        s[field] = "NOT AVAILABLE"
        s[rf]    = True

    # Run Setter I then II in the SAME loop so yellow counter stays accurate
    for s in subjects:
        if s["subject_type"] == "Lab":
            s["setter1_na"] = s["setter2_na"] = True
            continue
        assign_setter(s, "setter1", "setter1_yellow", "setter1_red", "setter1_na")
        s1 = s["setter1"] if not s["setter1_red"] else None
        assign_setter(s, "setter2", "setter2_yellow", "setter2_red", "setter2_na", exclude=s1)

    # ── Build per-subject setter exclusion set ────────────────────────────────
    # CRITICAL: any faculty assigned as Setter I or II (even if the OTHER setter
    # is "NOT AVAILABLE") must NEVER appear in PO / Examiner / Reserve for that row.
    def setter_excl(s):
        e = set()
        if s["setter1"] and not s["setter1_red"] and not s["setter1_na"]:
            e.add(s["setter1"])
        if s["setter2"] and not s["setter2_red"] and not s["setter2_na"]:
            e.add(s["setter2"])
        return e

    # ── Practical Only + Reserve for Lab/Lab+Theory ───────────────────────────
    for s in subjects:
        lab_fac = list(s["lab_desc"]["Faculty name"])

        if s["subject_type"] == "Theory":
            s["practical_only"] = "not applicable"
            continue

        if s["subject_type"] == "Lab":
            # Reserve = coordinator if in lab list, else highest lab workload
            coord = s["coordinator"]
            if coord and coord in lab_fac: reserve = coord
            elif lab_fac:                  reserve = lab_fac[0]
            else:                          reserve = None
            s["reserve"] = reserve if reserve else "NOT AVAILABLE"

            po = [f for f in lab_fac if f != reserve]
            s["_po_list"]       = po
            s["practical_only"] = ", ".join(po) if po else "NOT AVAILABLE"

        else:  # Lab+Theory
            excl = setter_excl(s)

            # Reserve = highest Total Workload excluding setters
            reserve = next(
                (r["Faculty name"] for _, r in s["all_desc"].iterrows()
                 if r["Faculty name"] not in excl), None)
            s["reserve"] = reserve if reserve else "NOT AVAILABLE"

            # PO = lab faculty excluding setters and reserve
            # Setters are NEVER included in PO, even as a last resort.
            non_setter_lab = [f for f in lab_fac
                              if f not in excl and f != reserve]
            po = non_setter_lab   # empty list → "NOT AVAILABLE" below

            s["_po_list"]       = po
            s["practical_only"] = ", ".join(po) if po else "NOT AVAILABLE"

    # ── Examiner ──────────────────────────────────────────────────────────────
    for s in subjects:
        if s["subject_type"] == "Lab":
            s["examiner"]    = "not applicable"
            s["examiner_na"] = True
            continue

        excl    = setter_excl(s)
        reserve = s.get("reserve")
        # Only add reserve to skip if it's an actual faculty name, not the sentinel
        skip    = excl | ({reserve} if reserve and reserve != "NOT AVAILABLE" else set())

        if s["subject_type"] == "Theory":
            names = [r["Faculty name"] for _, r in s["all_desc"].iterrows()
                     if r["Faculty name"] not in skip]
            s["examiner"] = ", ".join(names) if names else "NOT AVAILABLE"

        else:  # Lab+Theory: PO list + remaining (excl setters + reserve)
            po_set    = set(s["_po_list"])
            remaining = [r["Faculty name"] for _, r in s["all_desc"].iterrows()
                         if r["Faculty name"] not in skip
                         and r["Faculty name"] not in po_set]
            names = s["_po_list"] + remaining
            s["examiner"] = ", ".join(names) if names else "NOT AVAILABLE"

    # ── Reserve for Theory ────────────────────────────────────────────────────
    # Reserve = highest workload excl. setters; rebuild Examiner excluding reserve.
    for s in subjects:
        if s["subject_type"] != "Theory": continue

        excl = setter_excl(s)
        reserve = next(
            (r["Faculty name"] for _, r in s["all_desc"].iterrows()
             if r["Faculty name"] not in excl), None)
        s["reserve"] = reserve if reserve else "NOT AVAILABLE"

        # Rebuild Examiner excluding reserve
        skip  = excl | ({reserve} if reserve else set())
        names = [r["Faculty name"] for _, r in s["all_desc"].iterrows()
                 if r["Faculty name"] not in skip]
        s["examiner"] = ", ".join(names) if names else "NOT AVAILABLE"

    # ── Build output rows ─────────────────────────────────────────────────────
    rows = []
    for i, s in enumerate(subjects, 1):
        rows.append({
            "Sr. No."        : i,
            "Stream"         : s["stream"],
            "Subject code"   : s["code"],
            "Subject Name"   : s["name"],
            "Course Type"    : s["course_type"],
            "Subject Type"   : s["subject_type"],
            "Semester"       : s["semester"],
            "Paper Setter I" : s["setter1"],
            "Paper Setter II": s["setter2"],
            "Practical Only" : s["practical_only"],
            "Examiner"       : s["examiner"],
            "Reserve"        : s["reserve"],
            # Colour flags
            "_s1_red"   : s["setter1_red"],
            "_s1_yellow": s["setter1_yellow"],
            "_s1_na"    : s["setter1_na"],
            "_s2_red"   : s["setter2_red"],
            "_s2_yellow": s["setter2_yellow"],
            "_s2_na"    : s["setter2_na"],
            "_ex_na"    : s["examiner_na"],
            "_po_na"    : s["subject_type"] == "Theory",
            "_re_red"   : s["reserve"] == "NOT AVAILABLE",
        })

    return rows, global_counter


# ── Faculty Counter builder ───────────────────────────────────────────────────

def build_faculty_counter(rows, global_counter):
    """
    For each faculty in global_counter, find every setter assignment:
    record stream, semester, subject code.  If same subject code appears
    in multiple streams within the same semester, append stream in parentheses.
    Returns list of dicts ready to write to the Faculty Counter sheet.
    """
    # Map: faculty -> set of (semester, stream, code) — deduped so a faculty
    # appearing as BOTH Setter I and Setter II in the same row is counted once.
    assignments = defaultdict(set)
    for r in rows:
        sem    = r["Semester"]
        stream = r["Stream"]
        code   = r["Subject code"]
        for role in ["Paper Setter I", "Paper Setter II"]:
            name = r.get(role)
            if name and name != "NOT AVAILABLE":
                assignments[name].add((sem, stream, code))

    # All unique semesters (sorted)
    all_sems = sorted({sem for r in rows for sem in [r["Semester"]]})

    records = []
    for name in sorted(global_counter.keys()):
        unique_codes = global_counter[name]   # set of unique subject codes
        rec = {"Faculty Name": name, "Repetition": len(unique_codes)}
        # Per semester: collect codes; if code appears in multiple streams
        # in same semester, append stream tag e.g. "2CEIT402(CE-AI),2CEIT402(CE-IT)"
        for sem in all_sems:
            sem_assignments = [(st, cd) for (s, st, cd) in assignments[name] if s == sem]
            if not sem_assignments:
                rec[f"Sem {sem}"] = ""
            else:
                # Group by code to detect multi-stream
                code_streams = defaultdict(list)
                for st, cd in sem_assignments:
                    code_streams[cd].append(st)
                parts = []
                for cd, streams in code_streams.items():
                    if len(streams) == 1:
                        parts.append(cd)
                    else:
                        parts.append(",".join(f"{cd}({st})" for st in streams))
                rec[f"Sem {sem}"] = ", ".join(parts)
        records.append(rec)

    return records, all_sems


# ── Excel output columns ──────────────────────────────────────────────────────

COLUMNS = [
    "Sr. No.", "Stream", "Subject code", "Subject Name",
    "Course Type", "Subject Type",
    "Paper Setter I", "Paper Setter II",
    "Practical Only", "Examiner", "Reserve",
]

COL_WIDTHS = {
    "Sr. No.": 7, "Stream": 10, "Subject code": 14, "Subject Name": 28,
    "Course Type": 12, "Subject Type": 14,
    "Paper Setter I": 22, "Paper Setter II": 22,
    "Practical Only": 38, "Examiner": 48, "Reserve": 24,
}


def write_excel(rows, global_counter):
    import openpyxl

    output = io.BytesIO()
    wb     = openpyxl.Workbook()
    wb.remove(wb.active)

    sem_groups = defaultdict(list)
    for r in rows:
        sem_groups[r["Semester"]].append(r)

    # Shared styles
    hf   = Font(bold=True, color="FFFFFF")
    hfil = PatternFill("solid", fgColor="1C1B19")
    ca   = Alignment(horizontal="center", vertical="center", wrap_text=True)
    la   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    ts   = Side(style="thin", color="CCCCCC")
    tb   = Border(left=ts, right=ts, top=ts, bottom=ts)

    s1c  = COLUMNS.index("Paper Setter I")  + 1
    s2c  = COLUMNS.index("Paper Setter II") + 1
    poc  = COLUMNS.index("Practical Only")  + 1
    exc  = COLUMNS.index("Examiner")        + 1
    rec  = COLUMNS.index("Reserve")         + 1

    # ── Semester sheets ───────────────────────────────────────────────────────
    for sem in sorted(sem_groups.keys()):
        ws = wb.create_sheet(title=f"Sem {sem}" if sem else "Unknown Sem")
        ws.row_dimensions[1].height = 28

        for ci, cn in enumerate(COLUMNS, 1):
            c = ws.cell(row=1, column=ci, value=cn)
            c.font = hf; c.fill = hfil; c.alignment = ca; c.border = tb

        for ri, r in enumerate(sem_groups[sem], 2):
            ws.row_dimensions[ri].height = 20
            for ci, cn in enumerate(COLUMNS, 1):
                val  = r.get(cn)
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.alignment = la if ci > 3 else ca
                cell.border    = tb

                # Paper Setter I
                if ci == s1c:
                    if r["_s1_na"]:
                        cell.value = "not applicable"
                    elif r["_s1_red"]:
                        cell.fill = RED_FILL
                    elif r["_s1_yellow"]:
                        cell.fill = YELLOW_FILL

                # Paper Setter II
                elif ci == s2c:
                    if r["_s2_na"]:
                        cell.value = "not applicable"
                    elif r["_s2_red"]:
                        cell.fill = RED_FILL
                    elif r["_s2_yellow"]:
                        cell.fill = YELLOW_FILL

                # Practical Only
                elif ci == poc and r["_po_na"]:
                    cell.value = "not applicable"

                # Examiner
                elif ci == exc:
                    if r["_ex_na"]:
                        cell.value = "not applicable"
                    elif val == "NOT AVAILABLE":
                        cell.fill = RED_FILL

                # Reserve
                elif ci == rec and r["_re_red"]:
                    cell.fill = RED_FILL

        # Legend
        lr = len(sem_groups[sem]) + 3
        ws.row_dimensions[lr].height = 20

        def leg(col, text, fill, fc="1C1B19"):
            c = ws.cell(row=lr, column=col, value=text)
            c.fill = fill; c.font = Font(bold=True, color=fc)
            c.alignment = ca; c.border = tb
            ws.merge_cells(start_row=lr, start_column=col,
                           end_row=lr,   end_column=col + 1)

        leg(1, "No setter / role not found",  RED_FILL,    "FFFFFF")
        leg(3, "Faculty repeated in setter",  YELLOW_FILL, "1C1B19")

        for ci, cn in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(ci)].width = COL_WIDTHS.get(cn, 18)

    # ── Faculty Counter sheet ─────────────────────────────────────────────────
    counter_records, all_sems = build_faculty_counter(rows, global_counter)

    fc_cols = ["Faculty Name", "Repetition"] + [f"Sem {s}" for s in all_sems]
    fc_widths = {"Faculty Name": 28, "Repetition": 12}
    for s in all_sems:
        fc_widths[f"Sem {s}"] = 36

    fc = wb.create_sheet(title="Faculty Counter")
    fc.row_dimensions[1].height = 28

    for ci, cn in enumerate(fc_cols, 1):
        c = fc.cell(row=1, column=ci, value=cn)
        c.font = hf; c.fill = hfil; c.alignment = ca; c.border = tb

    for ri, rec_dict in enumerate(counter_records, 2):
        fc.row_dimensions[ri].height = 18
        for ci, cn in enumerate(fc_cols, 1):
            val  = rec_dict.get(cn, "")
            cell = fc.cell(row=ri, column=ci, value=val)
            cell.alignment = la if ci != 2 else ca
            cell.border    = tb
            # Highlight repeated faculty yellow
            if cn == "Repetition" and isinstance(val, int) and val > 1:
                cell.fill = YELLOW_FILL
            if cn == "Faculty Name" and rec_dict.get("Repetition", 0) > 1:
                cell.fill = YELLOW_FILL

    for ci, cn in enumerate(fc_cols, 1):
        fc.column_dimensions[get_column_letter(ci)].width = fc_widths.get(cn, 30)

    wb.save(output)
    output.seek(0)
    return output


# ── Flask routes ──────────────────────────────────────────────────────────────

@app.route("/")
def index():
    return render_template_string(HTML, error=None)


@app.route("/generate", methods=["POST"])
def generate():
    f = request.files.get("file")
    if not f or f.filename == "":
        return render_template_string(HTML, error="No file selected.")
    if not f.filename.lower().endswith(".xlsx"):
        return render_template_string(HTML, error="Only .xlsx files are supported.")
    try:
        df = pd.read_excel(f)
        required = {"Stream", "Subject code", "Subject Name", "Is Coordinator",
                    "Theory Workload", "Lab Workload", "Total Workload", "Faculty name"}
        missing = required - set(df.columns)
        if missing:
            return render_template_string(
                HTML, error=f"Missing columns: {', '.join(sorted(missing))}")
        rows, counter = run_algorithm(df)
        output = write_excel(rows, counter)
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="faculty_assignments.xlsx")
    except Exception as e:
        import traceback
        return render_template_string(HTML, error=f"Error:\n{traceback.format_exc()}")


if __name__ == "__main__":
    print("\n  Running at: http://localhost:5000\n")
    app.run(debug=True, port=5000)
